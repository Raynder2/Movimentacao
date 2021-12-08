from tkinter import *
import clipboard as cb
import tkinter.font as font
import pyautogui as pg
from tkinter import messagebox
from datetime import date
import pandas as pd
from openpyxl import *
import time
import sys

import re

cont = 1

# Dados fixos
motivo = "Movimentacao nova estrutura"
dataInicio = "01/01/2021"

# Mensagem solicitando execução
if(messagebox.askquestion("Confirmação", "Iniciar automoção?",icon ='info') == 'yes'):

    # Carregando planilha com as informações
    pb = load_workbook(filename="servidores.xlsx")
    sh = pb['Sheet1']

    # Dados que iremos inserir
    mat = sh['A'+str(cont)].value
    codlotacao = sh['B'+str(cont)].value
    nomlotacao = sh['C'+str(cont)].value

    pg.click(x=1577, y=397)
    time.sleep(1)
    pg.hotkey('f12')
    time.sleep(.5)


    # Comencando o laco de repetição
    while(True):
        if(cont > 1):
            mat = sh['A'+str(cont)].value
            codlotacao = sh['B'+str(cont)].value
            nomlotacao = sh['C'+str(cont)].value
        if(mat > 100):
            pg.click(x=1577, y=397)
            pg.write("f2 = document.querySelector('frame[name=\"f2\"]');")
            pg.write("documento = f2.contentDocument;")
            pg.write("frame = documento.querySelector('frame[name=\"cpo\"]');")
            pg.write("frame = frame.contentDocument;")

            time.sleep(3)
            pg.hotkey('enter')

            time.sleep(1)
            pg.write("frame.querySelector('input[name=\"txt_nr_matr_sif\"]').value = '"+str(mat)+"';")
            pg.hotkey('enter')

            time.sleep(2)
            pg.write("as = frame.querySelectorAll('a');")
            pg.write("as[0].click();")
            pg.hotkey('enter')

            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('enter')
            time.sleep(.5)

            pg.write("frame.querySelector('select[name=\"txt_tp_hist\"]').value = 4;")
            pg.press('enter')
            time.sleep(.5)
            pg.click(x=1577, y=397)
            
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('tab')
            pg.press('down')

            time.sleep(.5)
            pg.click(x=1577, y=397)
            pg.write("frame.querySelector('input[name=\"txt_cd_lot_novo\"]').value = '"+str(codlotacao)+"';")
            pg.write("frame.querySelector('input[name=\"txt_nm_lot_novo\"]').value = '"+str(nomlotacao)+"';")
            pg.write("frame.querySelector('input[name=\"txt_ds_mot_hs_sif\"]').value = 'CORRECAO DE LOTACAO';")
            pg.write("frame.querySelector('input[name=\"txt_dt_ini_freq_novo\"]').value = '15/09/2021';") #Definir inicio
            pg.write("frame.querySelector('input[name=\"txt_dt_fim_freq\"]').value = '14/09/2021';") #De
            
            time.sleep(3)
            pg.click(x=1577, y=397)
            pg.press('enter')
            pg.write("frame.querySelector('input[name=\"btn_incluir\"]').click();")
            
            pg.press('enter')
            time.sleep(1)

            pg.write("f2 = document.querySelector('frame[name=\"f2\"]');")
            pg.write("documento = f2.contentDocument;")
            pg.write("frame = documento.querySelector('frame[name=\"mnu\"]');")
            pg.write("frame = frame.contentDocument;")

            time.sleep(3)
            pg.hotkey('enter')

            pg.write("resposta = frame.querySelector('input[name=\"msg_n\"]').value;")
            pg.press('enter')
            time.sleep(1)
            pg.write("window.open(resposta, '_blank');")
            pg.press('enter')
            time.sleep(1)
            pg.hotkey("ctrl","l")
            time.sleep(2)
            pg.hotkey("ctrl","x")
            resposta = cb.paste()
            print(resposta)
            
            resul = len(resposta.split("SUCESS"))
            pg.hotkey("ctrl","w")

            if(resul > 1):
                sh['D'+str(cont)].value = "MOVIMENTADO"
                pb.save(filename="servidores.xlsx")
                pg.write("frame.querySelector('input[name=\"btn_limpar\"]').click();")
                pg.press('enter')
                time.sleep(1)
                pg.hotkey("ctrl","l")
            else:
                sh['D'+str(cont)].value = "Erro"
                pb.save(filename="servidores.xlsx")
                messagebox.showwarning(title='Erro', message='Nao foi possivel alterar a matricula '+str(mat))
                break
        else:
            break
        cont = cont + 1

    # Finalizando laço de repetição



    # pg.write("*/frame.querySelector('input[name=\"txt_cd_lot\"]').value")

# def alterar_linha(arquivo,nome):
#     cont = 0
#     with open(arquivo,'r') as f:
#         texto=f.readlines()
#     with open(arquivo,'w') as f:
#         exist = 0
#         for i in texto:
#             cont = cont + 1
#             if(cont == 11):
#                 f.write(i)
#                 messagebox.showwarning("Error","limite de atestados exedido, favor envialos a junta e somente apos enviados, adicionar novos.")
#                 return True
#             if(str(nome.strip()) in str(texto)):
#                 exist = 1
#             f.write(i)
#         messagebox.showinfo("Registrado", "Ok, pode guardar o atestado.")
#         if(exist == 1):
#             return True
#         f.write(str(nome)+'\n')
#         return True

# def cadastrar():
#     nome = entry.get()
#     alterar_linha('bd.txt', nome)
#     entry.delete('0', END)

# def limpar():
#     open('bd.txt', 'w').close()

# lista = ['POS1', 'POSIC1', 'POS3', 'POS4', 'POS5', 'POS6', 'POS7', 'POS8', 'POS9', 'POSI0', 'POSI1', 'POSIC2']
# doc = Document('K:/Administrativo/SetorPessoal/Raynder/Bot/arquivos_bases/junta.docx')

# def replaceWord(word, replace):
#     num = entry.get()
#     for paragraph in doc.paragraphs:
#         if word in paragraph.text:
#             inline = paragraph.runs
#             for i in range(len(inline)):
#                 if word in inline[i].text:
#                     text = inline[i].text.replace(word, replace)
#                     inline[i].text = text 
#     strcod = str(codigo)
#     strdia = str(dia)
#     strmes = str(mes)
#     strano = str(ano)
#     doc.save("K:/Administrativo/SetorPessoal/Raynder/Atestados/Remessas/relatorio"+strcod+"Junta"+strdia+"-"+strmes+"-"+strano+".doc")

# def gerarRelatorio():
#     novocodigo = int(codigo)+1
#     arquivo2 = open('K:/Administrativo/SetorPessoal/Raynder/Atestados/RegistroAtestados/numAtestado.txt', 'w')
#     arquivo2.write(str(novocodigo)+'\n')

#     contger = 0
#     with open('bd.txt','r') as f:
#         texto=f.readlines()
#     with open('bd.txt','w') as f:
#         for i in texto:
#             replaceWord(lista[contger], i.strip())
#             contger = contger + 1
#         replaceWord('DIA', str(dia))
#         replaceWord('MZ', str(meses[mes]))
#         replaceWord('AZ', str(ano))
#         replaceWord('RZ', str(codigo))

#         for c in lista:
#             replaceWord(c, "")

# myFont = font.Font(size=30)
# myFont2 = font.Font(size=15)

# C = Canvas(app, bg="blue", height=250, width=300)
# filename = PhotoImage(file = "Paisagem.png")
# background_label = Label(app, image=filename)
# background_label.place(x=0, y=0, relwidth=1, relheight=1)

# font1 = font.Font(name='TkCaptionFont', exists=True)
# font1.config(family='courier new', size=20)

# entry = AutocompleteEntry(autocompleteList, app, listboxLength=6, width=32, matchesFunction=matches)
# entry.place(x=150, y=45, width=250, height=20)
# button3 = Button(text='Cadastrar', bg='#145D3B', fg='#ffffff' , command=cadastrar)#bg fundo fg font
# button3['font'] = myFont
# button3.place(x=165, y=70, width=200, height=70)


# button3 = Button(text='Gerar Documento', bg='#FF0000', fg='#ffffff' , command=gerarRelatorio)#bg fundo fg font
# button3['font'] = myFont2
# button3.place(x=410, y=340, width=170, height=35)

# app.mainloop()